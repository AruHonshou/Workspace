from __future__ import annotations

import re
from difflib import SequenceMatcher
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from pypdf import PdfReader
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .ai_contracts import ATSResumeProposal
from .career import cloud_safe_verified_facts, detect_job_language, technology_mentions
from .documents import register_unicode_document_fonts
from .schemas import (
    ATSResumeDocument,
    ATSResumeLine,
    JobRecord,
    PrivateContactBlock,
    Profile,
)

_VISIBLE_REFERENCE_RE = re.compile(
    r"\s*\[?\s*(?:fact|record)[ _-]?id\s*[:=]\s*[A-Za-z0-9_-]+\s*\]?\s*",
    re.IGNORECASE,
)


def protected_tokens(value: str) -> set[str]:
    """Return facts a rewrite must preserve without freezing ordinary wording."""

    tokens = re.findall(r"\b[\w.+#%-]+\b", value, flags=re.UNICODE)
    protected: set[str] = set()
    for index, token in enumerate(tokens):
        if (
            re.search(r"\d", token)
            or (len(token) > 1 and token.isupper())
            or any(marker in token for marker in (".", "+", "#"))
            or (index > 0 and token[:1].isupper())
        ):
            protected.add(token.casefold())
    return protected


def _safe_proposed_text(original: str, proposed: str) -> str:
    """Reject rewrites that silently remove metrics, dates, or named terms."""

    cleaned = " ".join(_VISIBLE_REFERENCE_RE.sub(" ", proposed).split())
    if not cleaned:
        return original
    if not protected_tokens(original).issubset(protected_tokens(cleaned)):
        return original
    return cleaned


def _safe_context_heading(original: str, proposed: str | None) -> str | None:
    """Allow chronology labels only when their named details exist in evidence."""

    if not proposed:
        return None
    cleaned = " ".join(_VISIBLE_REFERENCE_RE.sub(" ", proposed).split())
    if not cleaned:
        return None
    if not protected_tokens(cleaned).issubset(protected_tokens(original)):
        return None
    return cleaned


def ats_adaptation_issues(
    document: ATSResumeDocument,
    job: JobRecord | None = None,
) -> list[str]:
    """Detect a model response that merely republishes the source résumé."""

    lines = [*document.experience, *document.projects]
    if not lines:
        return [
            "The proposal did not select any relevant experience or project evidence."
        ]
    similarities = [
        SequenceMatcher(
            None,
            " ".join(line.original_text.casefold().split()),
            " ".join(line.proposed_text.casefold().split()),
        ).ratio()
        for line in lines
    ]
    unchanged = sum(score >= 0.97 for score in similarities)
    issues: list[str] = []
    if unchanged == len(lines):
        issues.append(
            "The proposal copied every selected résumé line instead of tailoring it to the vacancy."
        )
    elif unchanged / len(lines) > 0.5:
        issues.append(
            "Most selected résumé lines remain materially unchanged and need stronger vacancy-specific wording."
        )
    if job is not None:
        ignored = {
            "and", "con", "de", "del", "el", "en", "engineer", "for", "la",
            "of", "para", "the", "un", "una", "y",
        }
        title_tokens = {
            token.casefold()
            for token in re.findall(r"[\w+#.-]+", job.title)
            if len(token) > 1 and token.casefold() not in ignored
        }
        visible_text = " ".join(
            [
                document.headline,
                document.professional_summary,
                *document.skills,
                *(line.proposed_text for line in lines),
            ]
        )
        visible_tokens = {
            token.casefold() for token in re.findall(r"[\w+#.-]+", visible_text)
        }
        if title_tokens and not title_tokens & visible_tokens:
            issues.append(
                "The headline and summary are not clearly targeted to the vacancy role."
            )

        job_technologies = technology_mentions(f"{job.title}\n{job.description}")
        evidenced_technologies = technology_mentions(
            "\n".join(line.original_text for line in lines)
        )
        supported_technologies = job_technologies & evidenced_technologies
        visible_technologies = technology_mentions(visible_text)
        if supported_technologies and not supported_technologies & visible_technologies:
            issues.append(
                "The proposal omitted technology keywords supported by the selected evidence."
            )
    return issues


def _fallback_document(profile: Profile, job: JobRecord) -> ATSResumeDocument:
    language = detect_job_language(job)
    facts = cloud_safe_verified_facts(
        profile,
        limit=max(1, len(profile.facts)),
        language=language,
        include_all=True,
    )
    experience: list[ATSResumeLine] = []
    education: list[ATSResumeLine] = []
    certifications: list[ATSResumeLine] = []
    languages: list[ATSResumeLine] = []
    skills: list[str] = []
    for fact in facts:
        line = ATSResumeLine(
            record_ids=[fact["fact_id"]],
            original_text=fact["text"],
            proposed_text=fact["text"],
        )
        category = fact.get("category", "")
        if category == "education":
            education.append(line)
        elif category == "certification":
            certifications.append(line)
        elif category == "language":
            languages.append(line)
        elif category == "skill":
            skills.append(fact["text"])
        else:
            experience.append(line)
    spanish = language.startswith("es")
    summary = (
        "Perfil profesional basado exclusivamente en experiencia y capacidades confirmadas."
        if spanish
        else "Professional profile based exclusively on confirmed experience and capabilities."
    )
    return ATSResumeDocument(
        language=language,
        headline=job.title,
        professional_summary=summary,
        skills=list(dict.fromkeys(skills))[:24],
        experience=experience[:24],
        education=education[:8],
        certifications=certifications[:8],
        languages=languages[:8],
    )


def build_ats_document(
    profile: Profile,
    job: JobRecord,
    proposal: ATSResumeProposal | None,
) -> tuple[ATSResumeDocument, list[str]]:
    """Validate a model proposal against the exact evidence offered to it."""

    fallback = _fallback_document(profile, job)
    if proposal is None:
        return fallback, [
            "The model proposal was unavailable; review the evidence-only draft."
        ]
    offered = {
        item["fact_id"]: item["text"]
        for item in cloud_safe_verified_facts(
            profile,
            limit=max(1, len(profile.facts)),
            language=detect_job_language(job),
            include_all=True,
        )
    }
    issues: list[str] = []

    def lines(values: list[object]) -> list[ATSResumeLine]:
        accepted: list[ATSResumeLine] = []
        for item in values:
            record_ids = list(getattr(item, "record_ids", []))
            if not record_ids or not set(record_ids).issubset(offered):
                issues.append(
                    "A proposed line referenced evidence that was not offered."
                )
                continue
            original = " ".join(offered[record_id] for record_id in record_ids)
            supplied_original = str(getattr(item, "original_text", ""))
            if (
                supplied_original
                and supplied_original not in original
                and original not in supplied_original
            ):
                issues.append("A proposed line changed its claimed source text.")
            proposed_text = _safe_proposed_text(
                original,
                str(getattr(item, "proposed_text", "")),
            )
            if proposed_text == original and proposed_text != str(
                getattr(item, "proposed_text", "")
            ):
                issues.append("A rewrite that removed protected facts was restored.")
            accepted.append(
                ATSResumeLine(
                    record_ids=record_ids,
                    original_text=original,
                    proposed_text=proposed_text,
                    context_heading=_safe_context_heading(
                        original,
                        getattr(item, "context_heading", None),
                    ),
                )
            )
        return accepted

    summary_record_ids = list(proposal.summary_record_ids)
    if not summary_record_ids or not set(summary_record_ids).issubset(offered):
        issues.append("The professional summary referenced unavailable evidence.")
        professional_summary = fallback.professional_summary
    else:
        professional_summary = " ".join(proposal.professional_summary.split())

    skills: list[str] = []
    for item in proposal.skills:
        record_ids = list(item.record_ids)
        if not record_ids or not set(record_ids).issubset(offered):
            issues.append("A proposed skill referenced unavailable evidence.")
            continue
        text = " ".join(_VISIBLE_REFERENCE_RE.sub(" ", item.text).split())
        evidence = " ".join(offered[record_id] for record_id in record_ids)
        evidence_tokens = {
            token.casefold() for token in re.findall(r"[\w.+#-]+", evidence)
        }
        candidates = [
            value.strip(" -•")
            for value in re.split(r"\s*(?:[,;•·]|\n)\s*", text)
            if value.strip(" -•")
        ] or [text]
        for candidate in candidates:
            if len(candidate) > 120:
                issues.append("A grouped skill was too long and was omitted.")
                continue
            skill_tokens = {
                token.casefold()
                for token in re.findall(r"[\w.+#-]+", candidate)
                if len(token) > 1
            }
            if not candidate or not skill_tokens or not skill_tokens & evidence_tokens:
                issues.append("A proposed skill was not found in its cited evidence.")
                continue
            skills.append(candidate)

    document = ATSResumeDocument(
        language=proposal.language,
        headline=" ".join(proposal.headline.split()) or fallback.headline,
        professional_summary=professional_summary,
        skills=list(dict.fromkeys(skills))[:30],
        experience=lines(proposal.experience),
        projects=lines(proposal.projects),
        education=lines(proposal.education),
        certifications=lines(proposal.certifications),
        languages=lines(proposal.languages),
    )
    issues.extend(ats_adaptation_issues(document, job))
    return document, list(dict.fromkeys(issues))


def _contact_lines(contact: PrivateContactBlock) -> list[str]:
    values = [
        *contact.emails[:1],
        *contact.phones[:1],
        *[str(item) for item in contact.websites[:2]],
    ]
    location = ", ".join(
        value for value in (contact.city, contact.region, contact.country_code) if value
    )
    if location:
        values.append(location)
    return values


def render_ats_docx(
    document: ATSResumeDocument,
    contact: PrivateContactBlock,
    output_path: Path,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    name = contact.full_name or document.headline
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(name)
    run.bold = True
    run.font.size = Pt(17)
    if contact.full_name:
        role = doc.add_paragraph(document.headline)
        role.alignment = WD_ALIGN_PARAGRAPH.CENTER
    details = _contact_lines(contact)
    if details:
        line = doc.add_paragraph(" | ".join(details))
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def heading(value: str) -> None:
        item = doc.add_paragraph()
        item.paragraph_format.space_before = Pt(8)
        item.paragraph_format.space_after = Pt(2)
        run = item.add_run(value.upper())
        run.bold = True
        run.font.size = Pt(11)

    def bullets(values: list[ATSResumeLine]) -> None:
        previous_heading: str | None = None
        for value in values:
            if value.context_heading and value.context_heading != previous_heading:
                context = doc.add_paragraph()
                context.paragraph_format.space_before = Pt(4)
                context.paragraph_format.space_after = Pt(1)
                context.add_run(value.context_heading).bold = True
                previous_heading = value.context_heading
            item = doc.add_paragraph(style="List Bullet")
            item.paragraph_format.space_after = Pt(1)
            item.add_run(value.proposed_text)

    labels = {
        "summary": "Resumen profesional"
        if document.language.startswith("es")
        else "Professional summary",
        "skills": "Habilidades" if document.language.startswith("es") else "Skills",
        "experience": "Experiencia"
        if document.language.startswith("es")
        else "Experience",
        "projects": "Proyectos" if document.language.startswith("es") else "Projects",
        "education": "Educación" if document.language.startswith("es") else "Education",
        "certifications": "Certificaciones"
        if document.language.startswith("es")
        else "Certifications",
        "languages": "Idiomas" if document.language.startswith("es") else "Languages",
    }
    heading(labels["summary"])
    doc.add_paragraph(document.professional_summary)
    if document.skills:
        heading(labels["skills"])
        doc.add_paragraph(" • ".join(document.skills))
    for key, values in (
        ("experience", document.experience),
        ("projects", document.projects),
        ("education", document.education),
        ("certifications", document.certifications),
        ("languages", document.languages),
    ):
        if values:
            heading(labels[key])
            bullets(values)
    doc.save(output_path)
    return output_path


def render_ats_pdf(
    document: ATSResumeDocument,
    contact: PrivateContactBlock,
    output_path: Path,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    font_regular, font_bold = register_unicode_document_fonts()
    normal = ParagraphStyle(
        "ATSNormal",
        parent=styles["BodyText"],
        fontName=font_regular,
        fontSize=9.5,
        leading=12,
        spaceAfter=4,
    )
    heading = ParagraphStyle(
        "ATSHeading",
        parent=normal,
        fontName=font_bold,
        fontSize=10.5,
        leading=13,
        spaceBefore=7,
        spaceAfter=3,
    )
    title = ParagraphStyle(
        "ATSTitle",
        parent=normal,
        fontName=font_bold,
        fontSize=16,
        leading=19,
        alignment=1,
    )
    story: list[object] = [
        Paragraph(escape(contact.full_name or document.headline), title)
    ]
    if contact.full_name:
        story.append(
            Paragraph(
                escape(document.headline),
                ParagraphStyle("Role", parent=normal, alignment=1),
            )
        )
    details = _contact_lines(contact)
    if details:
        story.append(
            Paragraph(
                escape(" | ".join(details)),
                ParagraphStyle("Contact", parent=normal, alignment=1),
            )
        )
    story.append(Spacer(1, 0.08 * inch))
    spanish = document.language.startswith("es")

    def section(label: str, values: list[str]) -> None:
        if not values:
            return
        story.append(Paragraph(escape(label.upper()), heading))
        for value in values:
            story.append(Paragraph(escape(value), normal))

    def evidence_section(label: str, values: list[ATSResumeLine]) -> None:
        if not values:
            return
        story.append(Paragraph(escape(label.upper()), heading))
        previous_heading: str | None = None
        for value in values:
            if value.context_heading and value.context_heading != previous_heading:
                story.append(
                    Paragraph(
                        escape(value.context_heading),
                        ParagraphStyle(
                            f"Context-{len(story)}",
                            parent=normal,
                            fontName=font_bold,
                            spaceBefore=3,
                            spaceAfter=1,
                        ),
                    )
                )
                previous_heading = value.context_heading
            story.append(Paragraph(f"• {escape(value.proposed_text)}", normal))

    section(
        "Resumen profesional" if spanish else "Professional summary",
        [document.professional_summary],
    )
    section(
        "Habilidades" if spanish else "Skills",
        [" • ".join(document.skills)] if document.skills else [],
    )
    evidence_section(
        "Experiencia" if spanish else "Experience",
        document.experience,
    )
    evidence_section(
        "Proyectos" if spanish else "Projects",
        document.projects,
    )
    evidence_section(
        "Educación" if spanish else "Education",
        document.education,
    )
    evidence_section(
        "Certificaciones" if spanish else "Certifications",
        document.certifications,
    )
    evidence_section(
        "Idiomas" if spanish else "Languages",
        document.languages,
    )
    pdf = SimpleDocTemplate(
        str(output_path),
        pagesize=LETTER,
        rightMargin=0.6 * inch,
        leftMargin=0.6 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
        title=document.headline,
        author="Workspace",
    )
    pdf.build(story)
    _validate_rendered_ats_pdf(output_path, document)
    return output_path


def _validate_rendered_ats_pdf(
    path: Path,
    document: ATSResumeDocument,
) -> None:
    """Fail closed when an ATS PDF is empty, corrupt, or leaks record IDs."""

    try:
        reader = PdfReader(path)
    except Exception as exc:  # pragma: no cover - parser exceptions vary by host.
        raise ValueError("Rendered ATS résumé is not a readable PDF") from exc
    if not 1 <= len(reader.pages) <= 6:
        raise ValueError(
            f"Rendered ATS résumé has an unexpected page count: {len(reader.pages)}"
        )
    page_texts = [(page.extract_text() or "").strip() for page in reader.pages]
    if any(len(text) < 30 for text in page_texts):
        raise ValueError("Rendered ATS résumé contains an empty or incomplete page")
    full_text = "\n".join(page_texts)
    if _VISIBLE_REFERENCE_RE.search(full_text):
        raise ValueError("Rendered ATS résumé exposes an internal record identifier")
    headline_words = re.findall(r"[a-z0-9]+", document.headline.casefold())[:3]
    extracted_words = " ".join(re.findall(r"[a-z0-9]+", full_text.casefold()))
    if headline_words and " ".join(headline_words) not in extracted_words:
        raise ValueError("Rendered ATS résumé is missing its headline")


__all__ = [
    "ats_adaptation_issues",
    "build_ats_document",
    "protected_tokens",
    "render_ats_docx",
    "render_ats_pdf",
]
